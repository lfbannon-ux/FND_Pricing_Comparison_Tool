#!/usr/bin/env python3
"""
Analytics & integrity-check blocks appended to each statement sheet, plus the
Key Metrics sheet.

Two hard rules from the brief are enforced here:
  * Growth is BLANK across any non-comparable boundary. Two periods are
    comparable only if they share a period_type and are ~one year apart.
    A stub period, a basis change, or a quarter-vs-year pairing yields blank.
  * Integrity checks are FORMULA rows (never hard values) so they re-evaluate
    on recalculation. They are labelled "CHECK:" so verify_workbook.py finds them.
"""
from __future__ import annotations

from datetime import date
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

F = "Arial"
NAVY = "1F3352"
LIGHT = "DCE3ED"
RULE = "8A97AA"
BLACK = "000000"
GREEN = "007A3D"
GREY = "5A6472"
PCT_CALC = "0.0%"
MONEY = '#,##0.0;(#,##0.0);"-"'
thin = Side(style="thin", color=RULE)


def _font(color=BLACK, bold=False, italic=False, sz=10):
    return Font(name=F, size=sz, bold=bold, italic=italic, color=color)


def _iso(s: str) -> date | None:
    try:
        y, m, d = (int(x) for x in s.split("-"))
        return date(y, m, d)
    except Exception:  # noqa: BLE001
        return None


def comparable(p_prev, p_cur) -> bool:
    """True only if a growth rate between these two periods is meaningful."""
    (_, e_prev, t_prev), (_, e_cur, t_cur) = p_prev, p_cur
    if t_prev != t_cur:
        return False
    d_prev, d_cur = _iso(e_prev), _iso(e_cur)
    if not d_prev or not d_cur:
        return False
    delta = (d_cur - d_prev).days
    return 330 <= delta <= 400


def find_row(line_rows: dict[str, int], *candidates: str) -> int | None:
    """Locate a canonical line by exact match, then by case-insensitive substring."""
    for cand in candidates:
        if cand in line_rows:
            return line_rows[cand]
    low = {k.lower(): v for k, v in line_rows.items()}
    for cand in candidates:
        c = cand.lower()
        if c in low:
            return low[c]
        for k, v in low.items():
            if c in k:
                return v
    return None


def _section(ws, r: int, text: str, ncols: int) -> int:
    for c in range(1, ncols + 2):
        ws.cell(row=r, column=c).fill = PatternFill("solid", fgColor=LIGHT)
    ws.cell(row=r, column=1, value=text).font = _font(NAVY, bold=True)
    return r + 1


def _ratio_row(ws, r, label, periods, num_row, den_row, ncols, fmt=PCT_CALC):
    ws.cell(row=r, column=1, value=label).font = _font(GREY)
    ws.cell(row=r, column=1).alignment = Alignment(indent=1)
    for i in range(len(periods)):
        col = get_column_letter(2 + i)
        cell = ws.cell(row=r, column=2 + i)
        cell.value = (f'=IF(OR(NOT(ISNUMBER({col}{den_row})),{col}{den_row}=0),"",'
                      f'{col}{num_row}/{col}{den_row})')
        cell.number_format = fmt
        cell.font = _font(BLACK)          # black = in-sheet formula
    return r + 1


def _growth_row(ws, r, label, periods, src_row, ncols):
    ws.cell(row=r, column=1, value=label).font = _font(GREY)
    ws.cell(row=r, column=1).alignment = Alignment(indent=1)
    for i in range(len(periods)):
        cell = ws.cell(row=r, column=2 + i)
        if i == 0 or not comparable(periods[i - 1], periods[i]):
            cell.value = None            # blank across a non-comparable boundary
        else:
            cur, prev = get_column_letter(2 + i), get_column_letter(1 + i)
            cell.value = (f'=IF(OR(NOT(ISNUMBER({prev}{src_row})),{prev}{src_row}=0),"",'
                          f'{cur}{src_row}/{prev}{src_row}-1)')
            cell.number_format = PCT_CALC
            cell.font = _font(BLACK)
    return r + 1


def _check_row(ws, r, label, periods, formula_fn, ncols):
    """A CHECK: row. verify_workbook.py asserts every value is ~0."""
    ws.cell(row=r, column=1, value=label).font = _font(GREY, italic=True)
    ws.cell(row=r, column=1).alignment = Alignment(indent=1)
    for i in range(len(periods)):
        col = get_column_letter(2 + i)
        cell = ws.cell(row=r, column=2 + i)
        cell.value = formula_fn(col)
        cell.number_format = MONEY
        cell.font = _font(BLACK)
        cell.border = Border(top=thin)
    return r + 1


def add_analytics_block(ws, r: int, periods: list, line_rows: dict[str, int],
                        statement: str, ncols: int,
                        cash_sheet_name: str | None = None,
                        cash_row: int | None = None,
                        bs_period_index: list | None = None) -> int:
    """Append the Analytics & checks block at the foot of a statement sheet."""
    r += 1
    r = _section(ws, r, "Analytics & checks", ncols)

    if statement == "income_statement":
        rev = find_row(line_rows, "Total revenues", "Revenue", "Revenues")
        op = find_row(line_rows, "Operating income", "Earnings from operations",
                      "Operating earnings")
        ni = find_row(line_rows, "Net earnings", "Net income", "Net earnings and comprehensive")
        if rev:
            r = _growth_row(ws, r, "Revenue growth", periods, rev, ncols)
        if rev and op:
            r = _ratio_row(ws, r, "Operating margin", periods, op, rev, ncols)
        if rev and ni:
            r = _ratio_row(ws, r, "Net margin", periods, ni, rev, ncols)
        if ni:
            r = _growth_row(ws, r, "Net earnings growth", periods, ni, ncols)

    elif statement == "balance_sheet":
        ta = find_row(line_rows, "Total assets")
        tl = find_row(line_rows, "Total liabilities")
        te = find_row(line_rows, "Total equity")
        tle = find_row(line_rows, "Total liabilities and equity")
        if ta and tl and te:
            r = _check_row(ws, r, "CHECK: total assets less total liabilities and equity",
                           periods, lambda c: f'=IF(ISNUMBER({c}{ta}),{c}{ta}-({c}{tl}+{c}{te}),"")',
                           ncols)
        if ta and tle:
            r = _check_row(ws, r, "CHECK: total assets less printed total liabilities and equity",
                           periods, lambda c: f'=IF(ISNUMBER({c}{tle}),{c}{ta}-{c}{tle},"")', ncols)
        ca = find_row(line_rows, "Total current assets")
        cl = find_row(line_rows, "Total current liabilities")
        if ca and cl:
            r = _ratio_row(ws, r, "Current ratio", periods, ca, cl, ncols, fmt='#,##0.00"x"')

    elif statement == "cash_flow":
        cfo = find_row(line_rows, "Cash from operating activities",
                       "Net cash provided by operating activities", "operating activities")
        cfi = find_row(line_rows, "Cash used in investing activities",
                       "Net cash used in investing activities", "investing activities")
        cff = find_row(line_rows, "Cash from financing activities",
                       "Net cash provided by financing activities", "financing activities")
        chg = find_row(line_rows, "Change in cash and cash equivalents",
                       "Net change in cash", "Increase (decrease) in cash")
        opening = find_row(line_rows, "Cash and cash equivalents, beginning of",
                           "beginning of period", "beginning of year")
        closing = find_row(line_rows, "Cash and cash equivalents, end of",
                           "end of period", "end of year")
        if cfo and cfi and cff and chg:
            r = _check_row(
                ws, r, "CHECK: operating + investing + financing less printed change in cash",
                periods,
                lambda c: f'=IF(ISNUMBER({c}{chg}),{c}{cfo}+{c}{cfi}+{c}{cff}-{c}{chg},"")', ncols)
        if opening and closing and chg:
            r = _check_row(
                ws, r, "CHECK: opening cash plus change less closing cash", periods,
                lambda c: f'=IF(ISNUMBER({c}{closing}),{c}{opening}+{c}{chg}-{c}{closing},"")',
                ncols)
        if closing and cash_sheet_name and cash_row:
            # GREEN: cross-sheet link to the balance sheet's cash line
            ws.cell(row=r, column=1,
                    value="CHECK: closing cash less balance-sheet cash (cross-sheet)"
                    ).font = _font(GREY, italic=True)
            ws.cell(row=r, column=1).alignment = Alignment(indent=1)
            ref = f"'{cash_sheet_name}'"
            for i in range(len(periods)):
                col = get_column_letter(2 + i)
                bs_col = get_column_letter(2 + bs_period_index[i]) if bs_period_index[i] is not None else None
                cell = ws.cell(row=r, column=2 + i)
                if bs_col is None:
                    cell.value = None            # period absent from the balance sheet
                else:
                    cell.value = (f'=IF(AND(ISNUMBER({col}{closing}),'
                                  f'ISNUMBER({ref}!{bs_col}{cash_row})),'
                                  f'{col}{closing}-{ref}!{bs_col}{cash_row},"")')
                    cell.number_format = MONEY
                    cell.font = _font(GREEN)
                cell.border = Border(top=thin)
            r += 1

    if statement in ('income_statement',):
        ws.cell(row=r, column=1,
                value="Growth is deliberately blank across any non-comparable boundary "
                      "(different period type, stub period, or basis change)."
                ).font = _font(GREY, italic=True, sz=9)
        r += 1
    return r
