"""The report and workbook writers must survive real data."""

import csv
import tempfile
import unittest
from pathlib import Path

from fnd_pricing import RETAILER_SHORT
from fnd_pricing.compare import compare_all
from fnd_pricing.loader import load_dataset
from fnd_pricing.report import render_html

ROOT = Path(__file__).resolve().parents[1]


class OutputTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        groups, offers = load_dataset(
            ROOT / "data/raw/sku_groups.csv", ROOT / "data/raw/products.csv"
        )
        cls.comparisons = compare_all(groups, offers)

    def test_html_report_is_self_contained(self):
        html = render_html(self.comparisons, "2026-08-14")
        self.assertIn("<!doctype html>", html)
        self.assertNotIn("http://", html.replace("http://www.w3.org", ""))
        self.assertNotIn("<script src=", html)
        self.assertIn("Luxury Vinyl Plank", html)

    def test_html_report_embeds_every_sku(self):
        html = render_html(self.comparisons)
        for comp in self.comparisons:
            self.assertIn(f'"{comp.group.group_id}"', html)

    def test_expense_section_is_rendered_with_both_charts(self):
        html = render_html(self.comparisons)
        self.assertIn("Expense by category", html)
        self.assertIn("Annual expense at Floor &amp; Decor prices", html)
        self.assertEqual(html.count('<svg class="chart"'), 2)
        self.assertIn("Below market low", html)   # legend, so colour is never alone

    def test_charts_carry_a_hover_and_keyboard_target_per_row(self):
        from fnd_pricing.spend import expense_by_category

        html = render_html(self.comparisons)
        categories = len(expense_by_category(self.comparisons))
        self.assertEqual(html.count('<g class="row" tabindex="0"'), categories * 2)

    def test_workbook_writes_all_sheets(self):
        from openpyxl import load_workbook

        from fnd_pricing.excel import write_workbook

        out = Path(tempfile.mkdtemp()) / "book.xlsx"
        write_workbook(self.comparisons, out)
        workbook = load_workbook(out)
        self.assertEqual(
            workbook.sheetnames,
            ["Summary", "Category Expense", "SKU Detail", "Offers", "Exceptions"],
        )
        self.assertEqual(workbook["SKU Detail"].max_row, len(self.comparisons) + 1)


if __name__ == "__main__":
    unittest.main()


class WorkbookShapeTest(unittest.TestCase):
    """Every sheet's header must cover every column its rows write.

    The per-competitor columns are generated, so a header list that falls out
    of step with the row builder produces a workbook with unlabelled columns
    rather than an error.
    """

    @classmethod
    def setUpClass(cls):
        import tempfile

        from openpyxl import load_workbook

        from fnd_pricing.excel import write_workbook

        groups, offers = load_dataset(
            ROOT / "data/raw/sku_groups.csv", ROOT / "data/raw/products.csv"
        )
        out = Path(tempfile.mkdtemp()) / "book.xlsx"
        write_workbook(compare_all(groups, offers), out)
        cls.workbook = load_workbook(out)

    def test_no_sheet_writes_past_its_header(self):
        for name in self.workbook.sheetnames:
            sheet = self.workbook[name]
            header_row = 1
            if name == "Summary":      # Summary's table starts below a stat block
                continue
            headers = [c.value for c in sheet[header_row]]
            self.assertTrue(
                all(h is not None for h in headers),
                f"{name}: {headers.count(None)} unlabelled column(s) - the header "
                f"list is out of step with the row builder",
            )

    def test_competitor_columns_scale_with_the_registry(self):
        from fnd_pricing import COMPETITORS

        headers = [c.value for c in self.workbook["SKU Detail"][1]]
        for retailer in COMPETITORS:
            short = RETAILER_SHORT.get(retailer, retailer)
            self.assertIn(f"{short} unit price", headers)


class RawPriceExportTest(unittest.TestCase):
    """The wide price export is the human-readable view of the canonical data."""

    @classmethod
    def setUpClass(cls):
        import tempfile

        from fnd_pricing.cli import main

        cls.path = Path(tempfile.mkdtemp()) / "prices.csv"
        main(["prices", "--csv", str(cls.path), "--show", "0"])
        with open(cls.path, encoding="utf-8") as handle:
            cls.rows = list(csv.DictReader(handle))

    def test_one_row_per_sku_group(self):
        groups, _ = load_dataset(
            ROOT / "data/raw/sku_groups.csv", ROOT / "data/raw/products.csv"
        )
        self.assertEqual(len(self.rows), len(groups))

    def test_every_retailer_gets_a_column_block(self):
        from fnd_pricing import RETAILER_CODES, RETAILERS

        header = self.rows[0].keys()
        for retailer in RETAILERS:
            code = RETAILER_CODES[retailer]
            for field in ("price_as_quoted", "uom", "unit_price", "match_tier"):
                self.assertIn(f"{code}_{field}", header)

    def test_a_missing_offer_is_blank_never_zero(self):
        # A zero in a price column would price the item as free and silently
        # win every comparison it appears in.
        blanks = 0
        for row in self.rows:
            for key, value in row.items():
                if key.endswith("_price_as_quoted") or key.endswith("_unit_price"):
                    self.assertNotEqual(value, "0")
                    self.assertNotEqual(value, "0.0")
                    if value == "":
                        blanks += 1
        self.assertGreater(blanks, 0, "expected some uncarried offers in the dataset")

    def test_quoted_price_and_unit_price_reconcile(self):
        from fnd_pricing import RETAILER_CODES, RETAILERS

        checked = 0
        for row in self.rows:
            for retailer in RETAILERS:
                code = RETAILER_CODES[retailer]
                quoted, unit = row[f"{code}_price_as_quoted"], row[f"{code}_unit_price"]
                pack, promo = row[f"{code}_pack_coverage"], row[f"{code}_promo_price"]
                if not quoted or not unit or promo:
                    continue
                effective = float(quoted) / float(pack) if pack else float(quoted)
                if row[f"{code}_uom"] == "per_sq_yd":
                    effective /= 9
                self.assertAlmostEqual(effective, float(unit), places=2)
                checked += 1
        self.assertGreater(checked, 100)


class EdlpSectionTest(unittest.TestCase):
    """The report's everyday-low-price read needs the study priced both ways."""

    @classmethod
    def setUpClass(cls):
        from fnd_pricing.compare import compare_all as _compare

        groups, offers = load_dataset(
            ROOT / "data/raw/sku_groups.csv", ROOT / "data/raw/products.csv"
        )
        cls.today = _compare(groups, offers)
        cls.shelf = _compare(groups, offers, use_list_price=True)
        cls.promo = {
            r: (sum(1 for o in offers if o.retailer == r and o.on_promo),
                sum(1 for o in offers if o.retailer == r), 0.15)
            for r in {o.retailer for o in offers}
        }

    def test_section_appears_only_when_both_views_are_supplied(self):
        with_both = render_html(self.today, "", self.shelf, self.promo)
        self.assertIn("Everyday low price", with_both)
        self.assertNotIn("Everyday low price", render_html(self.today))

    def test_the_chart_carries_a_market_low_row_beside_the_head_to_heads(self):
        from fnd_pricing import competitors

        html = render_html(self.today, "", self.shelf, self.promo)
        self.assertIn("Market low (best of all)", html)
        # One pair of bars per competitor, plus the market-low pair.
        self.assertEqual(
            html.count('class="mark series-a"'), len(competitors()) + 1
        )
        self.assertEqual(
            html.count('class="mark series-b"'), len(competitors()) + 1
        )

    def test_both_series_are_legended_so_colour_is_never_alone(self):
        html = render_html(self.today, "", self.shelf, self.promo)
        self.assertIn("As priced today", html)
        self.assertIn("At shelf price", html)
        self.assertIn('key series-a', html)
        self.assertIn('key series-b', html)

    def test_the_shelf_view_is_not_simply_the_same_numbers(self):
        # If promotions were being ignored on one side only, or not at all,
        # these two would agree; they must not.
        from fnd_pricing.compare import build_rollup

        today = build_rollup("t", self.today)
        shelf = build_rollup("s", self.shelf)
        self.assertNotEqual(today.win_rate, shelf.win_rate)
        self.assertGreater(shelf.win_rate, today.win_rate)
