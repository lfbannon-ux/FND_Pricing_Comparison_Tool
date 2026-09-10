"""Excel workbook export for merchants who want to slice the data themselves."""

from __future__ import annotations

from pathlib import Path
from typing import List

from openpyxl.utils import get_column_letter

from . import BASE_RETAILER, RETAILER_LABELS, RETAILER_SHORT, RETAILERS, competitors
from .compare import GroupComparison, build_rollup, rollup_by_category
from .spend import MARKET_LOW, expense_by_category, total_expense


_MONEY = '"$"#,##0.00'
_PCT = "+0.0%;-0.0%;0.0%"
_INDEX = "0.000"


def _style_header(worksheet, row: int = 1) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    fill = PatternFill("solid", fgColor="1F3B57")
    for cell in worksheet[row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    worksheet.freeze_panes = worksheet.cell(row=row + 1, column=1)


def _autosize(worksheet, limit: int = 46) -> None:

    for idx, column in enumerate(worksheet.columns, start=1):
        width = max((len(str(c.value)) for c in column if c.value is not None), default=8)
        worksheet.column_dimensions[get_column_letter(idx)].width = min(width + 2, limit)


def write_workbook(comparisons: List[GroupComparison], path: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    workbook = Workbook()
    overall = build_rollup("All categories", comparisons)

    # --- Summary ---
    sheet = workbook.active
    sheet.title = "Summary"
    sheet.append(["Floor & Decor like-for-like pricing comparison"])
    sheet["A1"].font = Font(bold=True, size=14)
    sheet.append([])
    for label, value, fmt in [
        ("SKU groups in study", overall.groups, "0"),
        ("SKU groups with a comparable competitor offer", overall.compared, "0"),
        ("Wins / ties / losses vs cheapest competitor",
         f"{overall.wins} / {overall.ties} / {overall.losses}", None),
        ("Win rate", overall.win_rate, "0.0%"),
        ("Median gap vs market low", overall.median_gap, _PCT),
        ("Mean gap vs market low", overall.mean_gap, _PCT),
        ("Volume-weighted basket index vs market low", overall.spend_index, _INDEX),
    ] + [
        (f"Median gap vs {RETAILER_LABELS.get(r, r)}", overall.per_retailer_gap.get(r), _PCT)
        for r in competitors()
    ]:
        sheet.append([label, value])
        if fmt:
            sheet.cell(row=sheet.max_row, column=2).number_format = fmt
        sheet.cell(row=sheet.max_row, column=1).font = Font(bold=True)

    sheet.append([])
    header_row = sheet.max_row + 1
    sheet.append([
        "Category", "SKUs", "Compared", "Win rate", "Median gap", "Mean gap",
        "Basket index",
    ] + [f"Median vs {RETAILER_SHORT.get(r, r)}" for r in competitors()])
    for rollup in rollup_by_category(comparisons):
        sheet.append([
            rollup.label, rollup.groups, rollup.compared, rollup.win_rate,
            rollup.median_gap, rollup.mean_gap, rollup.spend_index,
        ] + [rollup.per_retailer_gap.get(r) for r in competitors()])
        row = sheet.max_row
        sheet.cell(row=row, column=4).number_format = "0.0%"
        for col in (5, 6):
            sheet.cell(row=row, column=col).number_format = _PCT
        sheet.cell(row=row, column=7).number_format = _INDEX
        for offset in range(len(competitors())):
            sheet.cell(row=row, column=8 + offset).number_format = _PCT
    _style_header(sheet, header_row)
    sheet.freeze_panes = f"A{header_row + 1}"
    _autosize(sheet)

    # --- Category expense: the level, not just the rate ---
    expense = workbook.create_sheet("Category Expense")
    expense.append([
        "Category", "SKUs priced", "Basis", "Annual units", "Avg unit price",
        "Annual expense", "Share of expense", "Benchmarked expense",
        "Unbenchmarked expense", "Benchmark coverage", "Matched SKUs vs low",
        "Index vs market low", "$ vs market low",
    ] + [
        heading
        for r in competitors()
        for heading in (f"Index vs {RETAILER_SHORT.get(r, r)}",
                        f"$ vs {RETAILER_SHORT.get(r, r)}")
    ])
    rows = expense_by_category(comparisons)
    total = total_expense(comparisons)
    for row in rows + [total]:
        low = row.baskets[MARKET_LOW]
        expense.append([
            row.category, row.priced_skus, row.basis or "mixed", row.annual_units,
            row.avg_unit_price, row.annual_spend, row.share_of(total.annual_spend),
            row.benchmarked_spend, row.unbenchmarked_spend, row.benchmark_coverage,
            low.matched_skus, low.index, low.delta,
        ] + [
            value
            for r in competitors()
            for value in (row.baskets[r].index, row.baskets[r].delta)
        ])
        line = expense.max_row
        for col in (5, 6, 8, 9, 13):
            expense.cell(row=line, column=col).number_format = _MONEY
        for col in (7, 10):
            expense.cell(row=line, column=col).number_format = "0.0%"
        expense.cell(row=line, column=12).number_format = _INDEX
        # index / dollars alternate from column 14 onwards, one pair per competitor
        for offset in range(len(competitors())):
            expense.cell(row=line, column=14 + offset * 2).number_format = _INDEX
            expense.cell(row=line, column=15 + offset * 2).number_format = _MONEY
    for cell in expense[expense.max_row]:
        cell.font = Font(bold=True)
    _style_header(expense)
    expense.auto_filter.ref = (
        f"A1:{get_column_letter(expense.max_column)}{expense.max_row - 1}"
    )
    _autosize(expense)

    # --- SKU detail ---
    detail = workbook.create_sheet("SKU Detail")
    detail.append([
        "Group", "Category", "Subcategory", "Description", "Basis", "Annual volume",
        "F&D brand", "F&D unit price",
    ] + [
        heading
        for r in competitors()
        for heading in (
            f"{RETAILER_SHORT.get(r, r)} brand", f"{RETAILER_SHORT.get(r, r)} unit price",
            f"{RETAILER_SHORT.get(r, r)} delta %", f"{RETAILER_SHORT.get(r, r)} match",
            f"{RETAILER_SHORT.get(r, r)} counted",
        )
    ] + ["Market low", "Gap vs low", "Outcome", "Cheapest", "Flags"])
    for comp in comparisons:
        quotes = [comp.quotes[r] for r in competitors()]
        detail.append([
            comp.group.group_id, comp.group.category, comp.group.subcategory,
            comp.group.description, comp.group.basis, comp.group.annual_volume,
            comp.base.offer.brand if comp.base else "", comp.base_price,
        ] + [
            value
            for quote in quotes
            for value in (
                quote.normalized.offer.brand if quote.normalized else "",
                quote.unit_price, quote.delta_pct, quote.tier,
                "yes" if quote.included else "no",
            )
        ] + [
            comp.market_min, comp.gap_vs_market_min, comp.outcome,
            RETAILER_LABELS.get(comp.cheapest_retailer or "", ""),
            ", ".join(sorted({f.split(":")[0] for f in comp.flags})),
        ])
        row = detail.max_row
        detail.cell(row=row, column=8).number_format = _MONEY
        for offset in range(len(competitors())):
            base_col = 9 + offset * 5
            detail.cell(row=row, column=base_col + 1).number_format = _MONEY
            detail.cell(row=row, column=base_col + 2).number_format = _PCT
        tail = 9 + len(competitors()) * 5
        detail.cell(row=row, column=tail).number_format = _MONEY
        detail.cell(row=row, column=tail + 1).number_format = _PCT
    _style_header(detail)
    detail.auto_filter.ref = detail.dimensions
    _autosize(detail)

    # --- Offers: the audit trail behind every normalised price ---
    offers = workbook.create_sheet("Offers")
    offers.append([
        "Group", "Retailer", "Retailer SKU", "Brand", "Product", "List price", "UOM",
        "Pack coverage", "Promo price", "Effective price", "Normalised unit price",
        "Conversion", "In stock", "Collected on", "Data source", "URL",
    ])
    for comp in comparisons:
        normalized = ([comp.base] if comp.base else []) + [
            q.normalized for q in comp.quotes.values() if q.normalized
        ]
        for norm in normalized:
            offer = norm.offer
            offers.append([
                offer.group_id, RETAILER_LABELS.get(offer.retailer, offer.retailer),
                offer.retailer_sku, offer.brand, offer.product_name, offer.price,
                offer.uom, offer.pack_coverage, offer.promo_price, offer.effective_price,
                norm.unit_price, norm.conversion, "yes" if offer.in_stock else "no",
                offer.collected_on, offer.data_source, offer.url,
            ])
            row = offers.max_row
            for col in (6, 9, 10, 11):
                offers.cell(row=row, column=col).number_format = _MONEY
    _style_header(offers)
    offers.auto_filter.ref = offers.dimensions
    _autosize(offers)

    # --- Exceptions: what a human needs to look at before acting ---
    exceptions = workbook.create_sheet("Exceptions")
    exceptions.append(["Group", "Category", "Description", "Flag", "Detail"])
    for comp in comparisons:
        for flag in comp.flags:
            name, _, retailer = flag.partition(":")
            exceptions.append([
                comp.group.group_id, comp.group.category, comp.group.description, name,
                RETAILER_LABELS.get(retailer, retailer),
            ])
    _style_header(exceptions)
    exceptions.auto_filter.ref = exceptions.dimensions
    _autosize(exceptions)

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path
