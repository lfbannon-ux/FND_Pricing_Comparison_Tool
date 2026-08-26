#!/usr/bin/env python3
"""
Phase 6 -- verification.

1. Recalculate the finished workbook with headless LibreOffice.
2. Assert ZERO formula errors (#REF!/#DIV/0!/#NAME?/#VALUE!/#N/A/#NUM!).
3. Read the recalculated values back and confirm every integrity check ties.

Exit code 0 only if the workbook is clean.
"""
from __future__ import annotations

import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
BUILD = ROOT / "build"
SRC = BUILD / "Rockpoint_Gas_Storage_Historical_Model.xlsx"

ERRORS = ["#REF!", "#DIV/0!", "#NAME?", "#VALUE!", "#N/A", "#NUM!", "#NULL!", "Err:"]


def recalc(src: Path) -> Path:
    """Round-trip through LibreOffice so formulas carry cached values."""
    tmp = Path(tempfile.mkdtemp(prefix="rgsi_recalc_"))
    env_profile = tmp / "profile"
    cmd = [
        "soffice", "--headless", "--norestore", "--nolockcheck",
        f"-env:UserInstallation=file://{env_profile}",
        "--convert-to", "xlsx", "--outdir", str(tmp), str(src),
    ]
    res = subprocess.run(cmd, capture_output=True, text=True, timeout=600)
    out = tmp / src.name
    if not out.exists():
        print(res.stdout)
        print(res.stderr, file=sys.stderr)
        sys.exit("LibreOffice recalculation produced no output file.")
    return out


def main() -> int:
    if not SRC.exists():
        sys.exit(f"{SRC} not found -- run scripts/build_workbook.py first.")
    if not shutil.which("soffice"):
        sys.exit("soffice not on PATH -- cannot recalculate.")

    print(f"Recalculating {SRC.name} with headless LibreOffice ...")
    recalced = recalc(SRC)
    print(f"  -> {recalced}")

    wb = load_workbook(recalced, data_only=True)
    bad: list[str] = []
    cells_read = 0
    numeric = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if v is None:
                    continue
                cells_read += 1
                if isinstance(v, (int, float)):
                    numeric += 1
                elif isinstance(v, str):
                    for e in ERRORS:
                        if e in v:
                            bad.append(f"{ws.title}!{cell.coordinate} = {v}")
                            break

    print(f"\nCells with content : {cells_read:,}")
    print(f"Numeric cells      : {numeric:,}")
    print(f"Formula errors     : {len(bad)}")
    if bad:
        print("\n--- formula errors ---")
        for b in bad[:60]:
            print("  " + b)
        return 1

    # Integrity checks written by the builder are tagged in column A.
    checks = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            label = row[0].value
            if isinstance(label, str) and label.strip().lower().startswith("check:"):
                vals = [c.value for c in row[1:] if isinstance(c.value, (int, float))]
                worst = max((abs(v) for v in vals), default=0.0)
                checks.append((ws.title, label.strip(), worst, len(vals)))

    if checks:
        print("\n--- integrity checks (max |difference| across periods) ---")
        failed = 0
        for sheet, label, worst, n in checks:
            ok = worst < 0.05  # US$ millions, one decimal place presented
            status = "PASS" if ok else "FAIL"
            if not ok:
                failed += 1
            print(f"  [{status}] {sheet:<28} {label:<52} max={worst:,.4f} over {n} periods")
        if failed:
            print(f"\n{failed} integrity check(s) did not tie.")
            return 1
    else:
        print("\nNo integrity-check rows found (none written yet).")

    print("\nWorkbook is clean: zero formula errors"
          + (", all integrity checks tie." if checks else "."))
    return 0


if __name__ == "__main__":
    sys.exit(main())
